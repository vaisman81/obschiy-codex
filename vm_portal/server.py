from __future__ import annotations

import base64
import csv
import hashlib
import hmac
import json
import os
import secrets
import shutil
import sqlite3
import subprocess
import sys
import urllib.parse
import ipaddress
from datetime import datetime
from datetime import timedelta
from email.parser import BytesParser
from email.policy import default as email_default_policy
from http import HTTPStatus
from http.cookies import SimpleCookie
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


APP_DIR = Path(__file__).resolve().parent
DATA_DIR = APP_DIR / "data"
SNAPSHOT_DIR = DATA_DIR / "snapshots"
DB_PATH = DATA_DIR / "portal.sqlite3"
PUBLIC_DIR = APP_DIR / "public"
PBKDF2_ITERATIONS = 240_000
SESSION_IDLE_TIMEOUT_SECONDS = 10 * 60
IMPORT_TIMEOUT_SECONDS = 180


def now_iso() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def parse_iso(value: str) -> datetime:
    return datetime.strptime(value, "%Y-%m-%d %H:%M:%S")


def ensure_dirs() -> None:
    DATA_DIR.mkdir(exist_ok=True)
    SNAPSHOT_DIR.mkdir(exist_ok=True)


def connect_db() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA busy_timeout = 5000")
    return conn


def hash_password(password: str, salt: bytes | None = None) -> str:
    salt = salt or secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, PBKDF2_ITERATIONS)
    return f"pbkdf2_sha256${PBKDF2_ITERATIONS}${base64.b64encode(salt).decode()}${base64.b64encode(digest).decode()}"


def verify_password(password: str, stored: str) -> bool:
    try:
        algo, iterations, salt_b64, digest_b64 = stored.split("$", 3)
        if algo != "pbkdf2_sha256":
            return False
        salt = base64.b64decode(salt_b64)
        expected = base64.b64decode(digest_b64)
        actual = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, int(iterations))
        return hmac.compare_digest(actual, expected)
    except Exception:
        return False


def init_db() -> None:
    ensure_dirs()
    with connect_db() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE NOT NULL,
                password_hash TEXT NOT NULL,
                role TEXT NOT NULL CHECK(role IN ('read', 'write', 'admin')),
                created_at TEXT NOT NULL,
                is_locked INTEGER NOT NULL DEFAULT 0
            );
            CREATE TABLE IF NOT EXISTS sessions (
                token TEXT PRIMARY KEY,
                user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                created_at TEXT NOT NULL,
                last_seen TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS snapshots (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                created_at TEXT NOT NULL,
                uploaded_by INTEGER NOT NULL REFERENCES users(id),
                vsphere_filename TEXT NOT NULL,
                veeam_filename TEXT NOT NULL,
                total_vms INTEGER NOT NULL,
                powered_on INTEGER NOT NULL,
                powered_off INTEGER NOT NULL,
                has_backup INTEGER NOT NULL,
                has_replica INTEGER NOT NULL,
                off_with_backup INTEGER NOT NULL,
                off_with_replica INTEGER NOT NULL
            );
            CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS ip_bans (
                ip TEXT PRIMARY KEY,
                reason TEXT NOT NULL,
                banned_at TEXT NOT NULL,
                expires_at TEXT,
                banned_by INTEGER REFERENCES users(id)
            );
            CREATE TABLE IF NOT EXISTS auth_events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT,
                user_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
                ip TEXT NOT NULL,
                success INTEGER NOT NULL,
                reason TEXT NOT NULL,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS import_logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
                status TEXT NOT NULL,
                message TEXT NOT NULL,
                details TEXT,
                created_at TEXT NOT NULL
            );
            """
        )
        user_columns = {row["name"] for row in conn.execute("PRAGMA table_info(users)").fetchall()}
        if "is_locked" not in user_columns:
            conn.execute("ALTER TABLE users ADD COLUMN is_locked INTEGER NOT NULL DEFAULT 0")
        ban_columns = {row["name"] for row in conn.execute("PRAGMA table_info(ip_bans)").fetchall()}
        if "expires_at" not in ban_columns:
            conn.execute("ALTER TABLE ip_bans ADD COLUMN expires_at TEXT")
        admin_count = conn.execute("SELECT COUNT(*) FROM users WHERE role = 'admin'").fetchone()[0]
        if admin_count == 0:
            conn.execute(
                "INSERT INTO users(username, password_hash, role, created_at) VALUES (?, ?, ?, ?)",
                ("codex", hash_password(os.environ.get("VM_PORTAL_ADMIN_PASSWORD", "R8v#T4m!Q29x")), "admin", now_iso()),
            )
        session_columns = {row["name"] for row in conn.execute("PRAGMA table_info(sessions)").fetchall()}
        if "last_seen" not in session_columns:
            conn.execute("ALTER TABLE sessions ADD COLUMN last_seen TEXT")
            conn.execute("UPDATE sessions SET last_seen = created_at WHERE last_seen IS NULL")
        defaults = {
            "fail2ban_enabled": "true",
            "fail2ban_max_attempts": "5",
            "fail2ban_window_minutes": "10",
            "fail2ban_ban_minutes": "30",
        }
        for key, value in defaults.items():
            conn.execute("INSERT OR IGNORE INTO settings(key, value) VALUES (?, ?)", (key, value))


def read_csv_file(path: Path) -> list[dict[str, str]]:
    text = path.read_text(encoding="utf-8-sig")
    lines = [line for line in text.splitlines() if line.strip()]
    if lines and lines[0].startswith('"') and '""' in lines[0]:
        cleaned = []
        for line in lines:
            line = line.rstrip("\r\n").lstrip("\ufeff")
            if line.startswith('"'):
                line = line[1:]
            if line.endswith('"'):
                line = line[:-1]
            cleaned.append(line.replace('""', '"'))
        lines = cleaned
    reader = csv.DictReader(lines)
    return [{k: (v or "").strip() for k, v in row.items()} for row in reader]


def norm(value: str) -> str:
    return (value or "").strip().casefold()


def parse_veeam_date(value: str) -> datetime | None:
    for fmt in ("%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            pass
    return None


def analyze_pair(vsphere_path: Path, veeam_path: Path) -> tuple[list[dict], dict]:
    vsphere_rows = read_csv_file(vsphere_path)
    veeam_rows = read_csv_file(veeam_path)

    veeam_by_vm: dict[str, dict[str, dict]] = {}
    for row in veeam_rows:
        vm_name = row.get("VMName", "")
        protection = row.get("Protection", "")
        if not vm_name or protection not in {"Backup", "Replica"}:
            continue
        per_vm = veeam_by_vm.setdefault(
            norm(vm_name),
            {
                "Backup": {"jobs": set(), "count": 0, "latest": None, "latest_text": ""},
                "Replica": {"jobs": set(), "count": 0, "latest": None, "latest_text": ""},
            },
        )
        item = per_vm[protection]
        item["jobs"].add(row.get("JobName", ""))
        item["count"] += 1
        created = parse_veeam_date(row.get("CreationTime", ""))
        if created and (item["latest"] is None or created > item["latest"]):
            item["latest"] = created
            item["latest_text"] = row.get("CreationTime", "")

    records = []
    for row in vsphere_rows:
        vm_name = row.get("Name", "")
        if not vm_name or "replica" in norm(vm_name):
            continue
        vm_info = veeam_by_vm.get(norm(vm_name))
        backup = vm_info["Backup"] if vm_info else None
        replica = vm_info["Replica"] if vm_info else None
        backup_count = int(backup["count"]) if backup else 0
        replica_count = int(replica["count"]) if replica else 0
        records.append(
            {
                "vmName": vm_name,
                "powerState": row.get("PowerState", ""),
                "vmHost": row.get("VMHost", ""),
                "hasVeeamBackup": bool(backup_count),
                "hasVeeamReplica": bool(replica_count),
                "hasAnyVeeamProtection": bool(backup_count or replica_count),
                "backupJobs": sorted(job for job in (backup or {}).get("jobs", set()) if job),
                "replicaJobs": sorted(job for job in (replica or {}).get("jobs", set()) if job),
                "lastBackupRestorePoint": (backup or {}).get("latest_text", ""),
                "lastReplicaRestorePoint": (replica or {}).get("latest_text", ""),
                "backupRestorePoints": backup_count,
                "replicaRestorePoints": replica_count,
            }
        )

    records.sort(key=lambda item: (item["powerState"], item["vmName"].casefold()))
    summary = {
        "totalVms": len(records),
        "poweredOn": sum(1 for item in records if item["powerState"] == "PoweredOn"),
        "poweredOff": sum(1 for item in records if item["powerState"] == "PoweredOff"),
        "hasBackup": sum(1 for item in records if item["hasVeeamBackup"]),
        "hasReplica": sum(1 for item in records if item["hasVeeamReplica"]),
        "offWithBackup": sum(1 for item in records if item["powerState"] == "PoweredOff" and item["hasVeeamBackup"]),
        "offWithReplica": sum(1 for item in records if item["powerState"] == "PoweredOff" and item["hasVeeamReplica"]),
        "withoutProtection": sum(1 for item in records if not item["hasAnyVeeamProtection"]),
    }
    return records, summary


def write_report_xlsx(records: list[dict], snapshot: dict, output: BytesIO) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "VMs"
    headers = [
        "VMName",
        "PowerState",
        "VMHost",
        "HasVeeamBackup",
        "HasVeeamReplica",
        "HasAnyVeeamProtection",
        "BackupJobs",
        "ReplicaJobs",
        "LastBackupRestorePoint",
        "LastReplicaRestorePoint",
        "BackupRestorePoints",
        "ReplicaRestorePoints",
    ]
    ws.append(headers)
    for item in records:
        ws.append(
            [
                item["vmName"],
                item["powerState"],
                item["vmHost"],
                "Yes" if item["hasVeeamBackup"] else "No",
                "Yes" if item["hasVeeamReplica"] else "No",
                "Yes" if item["hasAnyVeeamProtection"] else "No",
                "; ".join(item["backupJobs"]),
                "; ".join(item["replicaJobs"]),
                item["lastBackupRestorePoint"],
                item["lastReplicaRestorePoint"],
                item["backupRestorePoints"] or "",
                item["replicaRestorePoints"] or "",
            ]
        )

    header_fill = PatternFill("solid", fgColor="1F2937")
    backup_fill = PatternFill("solid", fgColor="F8BBD0")
    replica_fill = PatternFill("solid", fgColor="D9EAD3")
    both_fill = PatternFill("solid", fgColor="D9B8FF")
    off_fill = PatternFill("solid", fgColor="F3F4F6")
    border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.border = border
    for row_idx in range(2, ws.max_row + 1):
        has_backup = ws.cell(row_idx, 4).value == "Yes"
        has_replica = ws.cell(row_idx, 5).value == "Yes"
        if ws.cell(row_idx, 2).value == "PoweredOff":
            ws.cell(row_idx, 2).fill = off_fill
        if has_backup and has_replica:
            fill = both_fill
            for col in (1, 4, 5):
                ws.cell(row_idx, col).fill = fill
        elif has_backup:
            for col in (1, 4):
                ws.cell(row_idx, col).fill = backup_fill
        elif has_replica:
            for col in (1, 5):
                ws.cell(row_idx, col).fill = replica_fill
        for col_idx in range(1, ws.max_column + 1):
            ws.cell(row_idx, col_idx).border = border
            ws.cell(row_idx, col_idx).alignment = Alignment(vertical="top", wrap_text=col_idx in {7, 8})

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in {
        "A": 42,
        "B": 14,
        "C": 18,
        "D": 16,
        "E": 16,
        "F": 20,
        "G": 32,
        "H": 32,
        "I": 22,
        "J": 22,
        "K": 18,
        "L": 18,
    }.items():
        ws.column_dimensions[col].width = width

    summary = wb.create_sheet("Summary")
    summary_rows = [
        ["Snapshot", snapshot.get("title", "")],
        ["Created", snapshot.get("created_at", "")],
        ["Total VMs", snapshot.get("total_vms", 0)],
        ["PoweredOn", snapshot.get("powered_on", 0)],
        ["PoweredOff", snapshot.get("powered_off", 0)],
        ["Has Backup", snapshot.get("has_backup", 0)],
        ["Has Replica", snapshot.get("has_replica", 0)],
        ["PoweredOff with Backup", snapshot.get("off_with_backup", 0)],
        ["PoweredOff with Replica", snapshot.get("off_with_replica", 0)],
    ]
    for row in summary_rows:
        summary.append(row)
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 36
    wb.save(output)


def json_response(handler: BaseHTTPRequestHandler, payload, status: int = 200) -> None:
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    handler.send_response(status)
    handler.send_header("Content-Type", "application/json; charset=utf-8")
    handler.send_header("Content-Length", str(len(body)))
    handler.end_headers()
    handler.wfile.write(body)


def parse_multipart_form(handler: BaseHTTPRequestHandler) -> tuple[dict[str, str], dict[str, tuple[str, bytes]]]:
    content_type = handler.headers.get("Content-Type", "")
    content_length = int(handler.headers.get("Content-Length", "0"))
    if "multipart/form-data" not in content_type or content_length <= 0:
        raise ValueError("Invalid multipart request")

    body = handler.rfile.read(content_length)
    raw = f"Content-Type: {content_type}\r\nMIME-Version: 1.0\r\n\r\n".encode("utf-8") + body
    message = BytesParser(policy=email_default_policy).parsebytes(raw)

    fields: dict[str, str] = {}
    files: dict[str, tuple[str, bytes]] = {}
    for part in message.iter_parts():
        name = part.get_param("name", header="content-disposition")
        if not name:
            continue
        filename = part.get_filename()
        payload = part.get_payload(decode=True) or b""
        if filename:
            files[name] = (Path(filename).name, payload)
        else:
            fields[name] = payload.decode("utf-8", errors="replace")
    return fields, files


def read_json(handler: BaseHTTPRequestHandler) -> dict:
    length = int(handler.headers.get("Content-Length", "0"))
    if length == 0:
        return {}
    return json.loads(handler.rfile.read(length).decode("utf-8"))


ROLE_LEVEL = {"read": 1, "write": 2, "admin": 3}


def client_ip(handler: BaseHTTPRequestHandler) -> str:
    forwarded = handler.headers.get("X-Forwarded-For", "")
    remote_ip = handler.client_address[0]
    if forwarded and remote_ip in {"127.0.0.1", "::1"}:
        return forwarded.split(",", 1)[0].strip()
    return remote_ip


def validate_ip_or_cidr(value: str) -> str | None:
    try:
        if "/" in value:
            return str(ipaddress.ip_network(value, strict=False))
        return str(ipaddress.ip_address(value))
    except ValueError:
        return None


def is_loopback_ip(value: str) -> bool:
    try:
        return ipaddress.ip_address(value).is_loopback
    except ValueError:
        return False


def get_settings(conn: sqlite3.Connection) -> dict[str, str]:
    return {row["key"]: row["value"] for row in conn.execute("SELECT key, value FROM settings").fetchall()}


def purge_expired_bans(conn: sqlite3.Connection) -> None:
    conn.execute("DELETE FROM ip_bans WHERE expires_at IS NOT NULL AND expires_at <= ?", (now_iso(),))


def is_ip_banned(conn: sqlite3.Connection, ip: str) -> bool:
    purge_expired_bans(conn)
    try:
        address = ipaddress.ip_address(ip)
    except ValueError:
        address = None
    for row in conn.execute("SELECT ip FROM ip_bans").fetchall():
        banned = row["ip"]
        if banned == ip:
            return True
        if address and "/" in banned:
            try:
                if address in ipaddress.ip_network(banned, strict=False):
                    return True
            except ValueError:
                continue
    return False


def is_last_active_admin(conn: sqlite3.Connection, user_id: int) -> bool:
    row = conn.execute("SELECT role, is_locked FROM users WHERE id = ?", (user_id,)).fetchone()
    if not row or row["role"] != "admin" or row["is_locked"]:
        return False
    active_admins = conn.execute("SELECT COUNT(*) FROM users WHERE role = 'admin' AND is_locked = 0").fetchone()[0]
    return active_admins <= 1


def safe_snapshot_dir(snapshot_id: str) -> Path | None:
    target = (SNAPSHOT_DIR / str(snapshot_id)).resolve()
    root = SNAPSHOT_DIR.resolve()
    if target == root or root not in target.parents:
        return None
    return target


def write_import_log(user_id: int | None, status: str, message: str, details: str = "") -> None:
    with connect_db() as conn:
        conn.execute(
            "INSERT INTO import_logs(user_id, status, message, details, created_at) VALUES (?, ?, ?, ?, ?)",
            (user_id, status, message[:300], details[-4000:], now_iso()),
        )


def record_auth_event(
    conn: sqlite3.Connection,
    username: str,
    user_id: int | None,
    ip: str,
    success: bool,
    reason: str,
) -> None:
    conn.execute(
        "INSERT INTO auth_events(username, user_id, ip, success, reason, created_at) VALUES (?, ?, ?, ?, ?, ?)",
        (username, user_id, ip, 1 if success else 0, reason, now_iso()),
    )


def apply_fail2ban(conn: sqlite3.Connection, ip: str) -> None:
    if is_loopback_ip(ip):
        return
    settings = get_settings(conn)
    if settings.get("fail2ban_enabled", "true") != "true":
        return
    max_attempts = int(settings.get("fail2ban_max_attempts", "5"))
    window_minutes = int(settings.get("fail2ban_window_minutes", "10"))
    ban_minutes = int(settings.get("fail2ban_ban_minutes", "30"))
    since = (datetime.now() - timedelta(minutes=window_minutes)).strftime("%Y-%m-%d %H:%M:%S")
    failures = conn.execute(
        """
        SELECT COUNT(*) FROM auth_events
        WHERE ip = ? AND success = 0 AND created_at >= ?
        """,
        (ip, since),
    ).fetchone()[0]
    if failures >= max_attempts:
        expires_at = (datetime.now() + timedelta(minutes=ban_minutes)).strftime("%Y-%m-%d %H:%M:%S")
        conn.execute(
            "INSERT OR REPLACE INTO ip_bans(ip, reason, banned_at, expires_at, banned_by) VALUES (?, ?, ?, ?, NULL)",
            (ip, f"fail2ban: {failures} failed logins in {window_minutes} minutes", now_iso(), expires_at),
        )


def powershell_executable() -> str:
    for candidate in ("pwsh.exe", "powershell.exe"):
        found = shutil.which(candidate)
        if found:
            return found
    return "powershell.exe"


def ps_quote(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def import_pair_from_servers(
    title: str,
    vmware_host: str,
    vmware_user: str,
    vmware_password: str,
    veeam_host: str,
    veeam_user: str,
    veeam_password: str,
    uploaded_by: int,
) -> tuple[int, dict]:
    tmp_id = secrets.token_hex(8)
    tmp_dir = SNAPSHOT_DIR / f"tmp-import-{tmp_id}"
    tmp_dir.mkdir(parents=True, exist_ok=True)
    vsphere_path = tmp_dir / "vsphere.csv"
    veeam_path = tmp_dir / "veeam.csv"
    script_path = tmp_dir / "collect.ps1"
    script = f"""
$ErrorActionPreference = 'Stop'
Set-StrictMode -Version Latest

$vmwareServer = {ps_quote(vmware_host)}
$vmwareUser = {ps_quote(vmware_user)}
$vmwarePassword = {ps_quote(vmware_password)}
$veeamServer = {ps_quote(veeam_host)}
$veeamUser = {ps_quote(veeam_user)}
$veeamPassword = {ps_quote(veeam_password)}
$vsphereCsv = {ps_quote(str(vsphere_path))}
$veeamCsv = {ps_quote(str(veeam_path))}

function Require-CommandOrThrow($Name, $InstallHint) {{
  if (-not (Get-Command $Name -ErrorAction SilentlyContinue)) {{
    throw "$Name is not available. $InstallHint"
  }}
}}

Import-Module VMware.PowerCLI -ErrorAction Stop
Set-PowerCLIConfiguration -InvalidCertificateAction Ignore -Confirm:$false | Out-Null
Require-CommandOrThrow 'Connect-VIServer' 'Install VMware PowerCLI in this Windows profile.'

if (Get-Module -ListAvailable -Name Veeam.Backup.PowerShell) {{
  Import-Module Veeam.Backup.PowerShell -ErrorAction Stop
}} elseif (Get-PSSnapin -Registered -Name VeeamPSSnapIn -ErrorAction SilentlyContinue) {{
  Add-PSSnapin VeeamPSSnapIn -ErrorAction Stop
}} else {{
  Write-Host "Local Veeam PowerShell not found. Using PowerShell remoting to $veeamServer."
}}

$vi = Connect-VIServer -Server $vmwareServer -User $vmwareUser -Password $vmwarePassword -ErrorAction Stop
Get-VM | Select-Object Name, PowerState, @{{Name='VMHost';Expression={{ $_.VMHost.Name }}}} |
  Export-Csv -NoTypeInformation -Encoding UTF8 -Path $vsphereCsv
Disconnect-VIServer -Server $vi -Confirm:$false | Out-Null

if (Get-Command Connect-VBRServer -ErrorAction SilentlyContinue) {{
  try {{
    Connect-VBRServer -Server $veeamServer -User $veeamUser -Password $veeamPassword -ErrorAction Stop | Out-Null
  }} catch {{
    Connect-VBRServer -Server $veeamServer -ErrorAction Stop | Out-Null
  }}
  $rows = New-Object System.Collections.Generic.List[object]
  Get-VBRBackup | ForEach-Object {{
    $backup = $_
    $protection = if (($backup.JobType -as [string]) -match 'Replica') {{ 'Replica' }} else {{ 'Backup' }}
    Get-VBRRestorePoint -Backup $backup | ForEach-Object {{
      $rows.Add([pscustomobject]@{{
        VMName = $_.Name
        Protection = $protection
        JobName = $backup.JobName
        CreationTime = $_.CreationTime.ToString('dd/MM/yyyy HH:mm:ss')
      }})
    }}
  }}
  $rows | Export-Csv -NoTypeInformation -Encoding UTF8 -Path $veeamCsv
  Disconnect-VBRServer | Out-Null
}} else {{
  $secure = ConvertTo-SecureString $veeamPassword -AsPlainText -Force
  $cred = New-Object System.Management.Automation.PSCredential($veeamUser, $secure)
  $session = New-PSSession -ComputerName $veeamServer -Credential $cred -ErrorAction Stop
  try {{
    $remoteCsv = Invoke-Command -Session $session -ScriptBlock {{
      $path = Join-Path $env:TEMP ("veeam-restore-points-" + [guid]::NewGuid().ToString("N") + ".csv")
      if (Get-Module -ListAvailable -Name Veeam.Backup.PowerShell) {{
        Import-Module Veeam.Backup.PowerShell -ErrorAction Stop
      }} elseif (Get-PSSnapin -Registered -Name VeeamPSSnapIn -ErrorAction SilentlyContinue) {{
        Add-PSSnapin VeeamPSSnapIn -ErrorAction Stop
      }} else {{
        throw "Veeam PowerShell is not installed on the Veeam server or is not available to this remote session."
      }}
      if (Get-Command Connect-VBRServer -ErrorAction SilentlyContinue) {{
        try {{ Connect-VBRServer -Server localhost -ErrorAction Stop | Out-Null }} catch {{ }}
      }}
      $rows = New-Object System.Collections.Generic.List[object]
      Get-VBRBackup | ForEach-Object {{
        $backup = $_
        $protection = if (($backup.JobType -as [string]) -match 'Replica') {{ 'Replica' }} else {{ 'Backup' }}
        Get-VBRRestorePoint -Backup $backup | ForEach-Object {{
          $rows.Add([pscustomobject]@{{
            VMName = $_.Name
            Protection = $protection
            JobName = $backup.JobName
            CreationTime = $_.CreationTime.ToString('dd/MM/yyyy HH:mm:ss')
          }})
        }}
      }}
      $rows | Export-Csv -NoTypeInformation -Encoding UTF8 -Path $path
      $path
    }} -ErrorAction Stop
    Copy-Item -FromSession $session -Path $remoteCsv -Destination $veeamCsv -Force -ErrorAction Stop
    Invoke-Command -Session $session -ScriptBlock {{ param($p) Remove-Item -LiteralPath $p -Force -ErrorAction SilentlyContinue }} -ArgumentList $remoteCsv | Out-Null
  }} finally {{
    Remove-PSSession $session
  }}
}}
"""
    script_path.write_text(script, encoding="utf-8")
    completed = subprocess.run(
        [powershell_executable(), "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script_path)],
        cwd=str(APP_DIR),
        capture_output=True,
        text=True,
        timeout=IMPORT_TIMEOUT_SECONDS,
    )
    script_path.unlink(missing_ok=True)
    if completed.returncode != 0:
        shutil.rmtree(tmp_dir, ignore_errors=True)
        details = (completed.stderr or completed.stdout or "PowerShell import failed").strip()
        raise RuntimeError(details[-1200:])
    if not vsphere_path.exists() or not veeam_path.exists():
        shutil.rmtree(tmp_dir, ignore_errors=True)
        raise RuntimeError("Import finished but CSV files were not created")

    records, summary = analyze_pair(vsphere_path, veeam_path)
    with connect_db() as conn:
        cur = conn.execute(
            """
            INSERT INTO snapshots(
                title, created_at, uploaded_by, vsphere_filename, veeam_filename,
                total_vms, powered_on, powered_off, has_backup, has_replica, off_with_backup, off_with_replica
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                title or f"Auto import {now_iso()}",
                now_iso(),
                uploaded_by,
                f"vmware-{vmware_host}.csv",
                f"veeam-{veeam_host}.csv",
                summary["totalVms"],
                summary["poweredOn"],
                summary["poweredOff"],
                summary["hasBackup"],
                summary["hasReplica"],
                summary["offWithBackup"],
                summary["offWithReplica"],
            ),
        )
        snapshot_id = cur.lastrowid

    final_dir = SNAPSHOT_DIR / str(snapshot_id)
    tmp_dir.rename(final_dir)
    (final_dir / "analysis.json").write_text(json.dumps({"records": records, "summary": summary}, ensure_ascii=False, indent=2), encoding="utf-8")
    return snapshot_id, summary


class PortalHandler(BaseHTTPRequestHandler):
    server_version = "SoftMasterVmPortal/0.1"

    def log_message(self, fmt: str, *args) -> None:
        print(f"{self.address_string()} - {fmt % args}")

    def current_user(self) -> sqlite3.Row | None:
        cookie = SimpleCookie(self.headers.get("Cookie"))
        token = cookie.get("session")
        if not token:
            return None
        with connect_db() as conn:
            session = conn.execute(
                """
                SELECT sessions.token, sessions.last_seen, users.id, users.username, users.role, users.is_locked
                FROM sessions JOIN users ON users.id = sessions.user_id
                WHERE sessions.token = ?
                """,
                (token.value,),
            ).fetchone()
            if not session:
                return None
            if session["is_locked"]:
                conn.execute("DELETE FROM sessions WHERE token = ?", (token.value,))
                return None
            try:
                last_seen = parse_iso(session["last_seen"])
            except Exception:
                last_seen = datetime.min
            if datetime.now() - last_seen > timedelta(seconds=SESSION_IDLE_TIMEOUT_SECONDS):
                conn.execute("DELETE FROM sessions WHERE token = ?", (token.value,))
                return None
            conn.execute("UPDATE sessions SET last_seen = ? WHERE token = ?", (now_iso(), token.value))
            return session

    def require_user(self, role: str = "read") -> sqlite3.Row | None:
        user = self.current_user()
        if not user:
            json_response(self, {"error": "Unauthorized"}, HTTPStatus.UNAUTHORIZED)
            return None
        if ROLE_LEVEL[user["role"]] < ROLE_LEVEL[role]:
            json_response(self, {"error": "Forbidden"}, HTTPStatus.FORBIDDEN)
            return None
        return user

    def do_GET(self) -> None:
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path
        if path.startswith("/api/"):
            return self.handle_api_get(path, parsed)
        return self.serve_static(path)

    def do_POST(self) -> None:
        parsed = urllib.parse.urlparse(self.path)
        if parsed.path.startswith("/api/"):
            return self.handle_api_post(parsed.path)
        json_response(self, {"error": "Not found"}, HTTPStatus.NOT_FOUND)

    def do_DELETE(self) -> None:
        parsed = urllib.parse.urlparse(self.path)
        if parsed.path.startswith("/api/"):
            return self.handle_api_delete(parsed.path)
        json_response(self, {"error": "Not found"}, HTTPStatus.NOT_FOUND)

    def serve_static(self, path: str) -> None:
        if path in {"", "/"}:
            path = "/index.html"
        target = (PUBLIC_DIR / path.lstrip("/")).resolve()
        if not str(target).startswith(str(PUBLIC_DIR.resolve())) or not target.exists() or target.is_dir():
            self.send_error(HTTPStatus.NOT_FOUND)
            return
        content_types = {
            ".html": "text/html; charset=utf-8",
            ".css": "text/css; charset=utf-8",
            ".js": "application/javascript; charset=utf-8",
            ".svg": "image/svg+xml",
        }
        body = target.read_bytes()
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", content_types.get(target.suffix, "application/octet-stream"))
        if target.suffix in {".html", ".css", ".js"}:
            self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def handle_api_get(self, path: str, parsed) -> None:
        if path == "/api/me":
            user = self.current_user()
            return json_response(self, {"user": dict(user) if user else None})
        if path == "/api/snapshots":
            if not self.require_user("read"):
                return
            with connect_db() as conn:
                rows = conn.execute("SELECT * FROM snapshots ORDER BY created_at DESC, id DESC").fetchall()
            return json_response(self, {"snapshots": [dict(row) for row in rows]})
        if path.startswith("/api/snapshots/") and path.endswith("/data"):
            if not self.require_user("read"):
                return
            snapshot_id = path.split("/")[3]
            return self.send_snapshot_data(snapshot_id)
        if path.startswith("/api/snapshots/") and path.endswith("/export"):
            if not self.require_user("read"):
                return
            snapshot_id = path.split("/")[3]
            return self.send_export(snapshot_id)
        if path == "/api/users":
            if not self.require_user("admin"):
                return
            with connect_db() as conn:
                rows = conn.execute("SELECT id, username, role, created_at, is_locked FROM users ORDER BY username").fetchall()
            return json_response(self, {"users": [dict(row) for row in rows]})
        if path == "/api/security":
            if not self.require_user("admin"):
                return
            with connect_db() as conn:
                purge_expired_bans(conn)
                settings = get_settings(conn)
                bans = conn.execute(
                    """
                    SELECT ip, reason, banned_at, expires_at, users.username AS banned_by_username
                    FROM ip_bans LEFT JOIN users ON users.id = ip_bans.banned_by
                    ORDER BY banned_at DESC
                    """
                ).fetchall()
                events = conn.execute(
                    """
                    SELECT auth_events.*, users.username AS resolved_username
                    FROM auth_events LEFT JOIN users ON users.id = auth_events.user_id
                    ORDER BY auth_events.id DESC LIMIT 200
                    """
                ).fetchall()
                ip_stats = conn.execute(
                    """
                    SELECT ip,
                           COUNT(*) AS total,
                           SUM(CASE WHEN success = 1 THEN 1 ELSE 0 END) AS success_count,
                           SUM(CASE WHEN success = 0 THEN 1 ELSE 0 END) AS failure_count,
                           MAX(created_at) AS last_seen
                    FROM auth_events
                    GROUP BY ip
                    ORDER BY last_seen DESC
                    """
                ).fetchall()
            return json_response(
                self,
                {
                    "settings": settings,
                    "bans": [dict(row) for row in bans],
                    "events": [dict(row) for row in events],
                    "ipStats": [dict(row) for row in ip_stats],
                },
            )
        if path == "/api/import/logs":
            if not self.require_user("write"):
                return
            with connect_db() as conn:
                rows = conn.execute(
                    """
                    SELECT import_logs.*, users.username
                    FROM import_logs LEFT JOIN users ON users.id = import_logs.user_id
                    ORDER BY import_logs.id DESC LIMIT 200
                    """
                ).fetchall()
            return json_response(self, {"logs": [dict(row) for row in rows]})
        json_response(self, {"error": "Not found"}, HTTPStatus.NOT_FOUND)

    def handle_api_post(self, path: str) -> None:
        if path == "/api/login":
            payload = read_json(self)
            username = payload.get("username", "")
            password = payload.get("password", "")
            ip = client_ip(self)
            with connect_db() as conn:
                if is_ip_banned(conn, ip):
                    record_auth_event(conn, username, None, ip, False, "ip_banned")
                    return json_response(self, {"error": "IP is blocked"}, HTTPStatus.FORBIDDEN)
                user = conn.execute("SELECT * FROM users WHERE lower(username) = lower(?)", (username,)).fetchone()
                if not user or not verify_password(password, user["password_hash"]):
                    record_auth_event(conn, username, user["id"] if user else None, ip, False, "invalid_credentials")
                    apply_fail2ban(conn, ip)
                    return json_response(self, {"error": "Invalid credentials"}, HTTPStatus.UNAUTHORIZED)
                if user["is_locked"]:
                    record_auth_event(conn, username, user["id"], ip, False, "user_locked")
                    return json_response(self, {"error": "User is locked"}, HTTPStatus.FORBIDDEN)
                token = secrets.token_urlsafe(32)
                timestamp = now_iso()
                conn.execute(
                    "INSERT INTO sessions(token, user_id, created_at, last_seen) VALUES (?, ?, ?, ?)",
                    (token, user["id"], timestamp, timestamp),
                )
                record_auth_event(conn, username, user["id"], ip, True, "login")
            body = json.dumps({"user": {"id": user["id"], "username": user["username"], "role": user["role"]}}, ensure_ascii=False).encode("utf-8")
            self.send_response(HTTPStatus.OK)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Set-Cookie", f"session={token}; HttpOnly; SameSite=Lax; Path=/")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        if path == "/api/logout":
            cookie = SimpleCookie(self.headers.get("Cookie"))
            token = cookie.get("session")
            if token:
                with connect_db() as conn:
                    conn.execute("DELETE FROM sessions WHERE token = ?", (token.value,))
            self.send_response(HTTPStatus.NO_CONTENT)
            self.send_header("Set-Cookie", "session=; Max-Age=0; Path=/")
            self.end_headers()
            return
        if path == "/api/upload":
            user = self.require_user("write")
            if not user:
                return
            return self.handle_upload(user)
        if path == "/api/users":
            if not self.require_user("admin"):
                return
            payload = read_json(self)
            username = payload.get("username", "").strip()
            password = payload.get("password", "")
            role = payload.get("role", "read")
            if not username or not password or role not in ROLE_LEVEL:
                return json_response(self, {"error": "Username, password and valid role are required"}, HTTPStatus.BAD_REQUEST)
            try:
                with connect_db() as conn:
                    conn.execute(
                        "INSERT INTO users(username, password_hash, role, created_at) VALUES (?, ?, ?, ?)",
                        (username, hash_password(password), role, now_iso()),
                    )
            except sqlite3.IntegrityError:
                return json_response(self, {"error": "User already exists"}, HTTPStatus.CONFLICT)
            return json_response(self, {"ok": True}, HTTPStatus.CREATED)
        if path == "/api/me/password":
            user = self.require_user("read")
            if not user:
                return
            payload = read_json(self)
            current_password = payload.get("currentPassword", "")
            new_password = payload.get("newPassword", "")
            if len(new_password) < 8:
                return json_response(self, {"error": "New password must be at least 8 characters"}, HTTPStatus.BAD_REQUEST)
            cookie = SimpleCookie(self.headers.get("Cookie"))
            token = cookie.get("session")
            with connect_db() as conn:
                row = conn.execute("SELECT password_hash FROM users WHERE id = ?", (user["id"],)).fetchone()
                if not row or not verify_password(current_password, row["password_hash"]):
                    return json_response(self, {"error": "Current password is incorrect"}, HTTPStatus.BAD_REQUEST)
                conn.execute("UPDATE users SET password_hash = ? WHERE id = ?", (hash_password(new_password), user["id"]))
                if token:
                    conn.execute("DELETE FROM sessions WHERE user_id = ? AND token <> ?", (user["id"], token.value))
            return json_response(self, {"ok": True})
        if path.startswith("/api/users/") and path.endswith("/password"):
            admin = self.require_user("admin")
            if not admin:
                return
            try:
                user_id = int(path.split("/")[3])
            except ValueError:
                return json_response(self, {"error": "Invalid user id"}, HTTPStatus.BAD_REQUEST)
            payload = read_json(self)
            new_password = payload.get("newPassword", "")
            if len(new_password) < 8:
                return json_response(self, {"error": "New password must be at least 8 characters"}, HTTPStatus.BAD_REQUEST)
            with connect_db() as conn:
                cur = conn.execute("UPDATE users SET password_hash = ? WHERE id = ?", (hash_password(new_password), user_id))
                if cur.rowcount == 0:
                    return json_response(self, {"error": "User not found"}, HTTPStatus.NOT_FOUND)
                conn.execute("DELETE FROM sessions WHERE user_id = ?", (user_id,))
            return json_response(self, {"ok": True})
        if path.startswith("/api/users/") and path.endswith("/role"):
            admin = self.require_user("admin")
            if not admin:
                return
            try:
                user_id = int(path.split("/")[3])
            except ValueError:
                return json_response(self, {"error": "Invalid user id"}, HTTPStatus.BAD_REQUEST)
            payload = read_json(self)
            role = payload.get("role", "")
            if role not in ROLE_LEVEL:
                return json_response(self, {"error": "Valid role is required"}, HTTPStatus.BAD_REQUEST)
            with connect_db() as conn:
                if role != "admin" and is_last_active_admin(conn, user_id):
                    return json_response(self, {"error": "Cannot demote the last active admin"}, HTTPStatus.BAD_REQUEST)
                cur = conn.execute("UPDATE users SET role = ? WHERE id = ?", (role, user_id))
                if cur.rowcount == 0:
                    return json_response(self, {"error": "User not found"}, HTTPStatus.NOT_FOUND)
                conn.execute("DELETE FROM sessions WHERE user_id = ?", (user_id,))
            return json_response(self, {"ok": True})
        if path.startswith("/api/users/") and path.endswith("/lock"):
            admin = self.require_user("admin")
            if not admin:
                return
            try:
                user_id = int(path.split("/")[3])
            except ValueError:
                return json_response(self, {"error": "Invalid user id"}, HTTPStatus.BAD_REQUEST)
            if user_id == admin["id"]:
                return json_response(self, {"error": "Cannot lock current user"}, HTTPStatus.BAD_REQUEST)
            with connect_db() as conn:
                if is_last_active_admin(conn, user_id):
                    return json_response(self, {"error": "Cannot lock the last active admin"}, HTTPStatus.BAD_REQUEST)
                cur = conn.execute("UPDATE users SET is_locked = 1 WHERE id = ?", (user_id,))
                if cur.rowcount == 0:
                    return json_response(self, {"error": "User not found"}, HTTPStatus.NOT_FOUND)
                conn.execute("DELETE FROM sessions WHERE user_id = ?", (user_id,))
            return json_response(self, {"ok": True})
        if path.startswith("/api/users/") and path.endswith("/unlock"):
            if not self.require_user("admin"):
                return
            try:
                user_id = int(path.split("/")[3])
            except ValueError:
                return json_response(self, {"error": "Invalid user id"}, HTTPStatus.BAD_REQUEST)
            with connect_db() as conn:
                cur = conn.execute("UPDATE users SET is_locked = 0 WHERE id = ?", (user_id,))
                if cur.rowcount == 0:
                    return json_response(self, {"error": "User not found"}, HTTPStatus.NOT_FOUND)
            return json_response(self, {"ok": True})
        if path == "/api/security/settings":
            if not self.require_user("admin"):
                return
            payload = read_json(self)
            enabled = "true" if payload.get("fail2ban_enabled") in {True, "true", "on", "1", 1} else "false"
            try:
                max_attempts = max(1, min(100, int(payload.get("fail2ban_max_attempts", 5))))
                window_minutes = max(1, min(1440, int(payload.get("fail2ban_window_minutes", 10))))
                ban_minutes = max(1, min(10080, int(payload.get("fail2ban_ban_minutes", 30))))
            except (TypeError, ValueError):
                return json_response(self, {"error": "Fail2ban settings must be valid numbers"}, HTTPStatus.BAD_REQUEST)
            with connect_db() as conn:
                conn.execute("INSERT OR REPLACE INTO settings(key, value) VALUES ('fail2ban_enabled', ?)", (enabled,))
                conn.execute("INSERT OR REPLACE INTO settings(key, value) VALUES ('fail2ban_max_attempts', ?)", (str(max_attempts),))
                conn.execute("INSERT OR REPLACE INTO settings(key, value) VALUES ('fail2ban_window_minutes', ?)", (str(window_minutes),))
                conn.execute("INSERT OR REPLACE INTO settings(key, value) VALUES ('fail2ban_ban_minutes', ?)", (str(ban_minutes),))
            return json_response(self, {"ok": True})
        if path == "/api/security/ban-ip":
            admin = self.require_user("admin")
            if not admin:
                return
            payload = read_json(self)
            ip = validate_ip_or_cidr(payload.get("ip", "").strip())
            reason = payload.get("reason", "manual ban").strip() or "manual ban"
            if not ip:
                return json_response(self, {"error": "Valid IP or CIDR is required"}, HTTPStatus.BAD_REQUEST)
            current_ip = client_ip(self)
            current_matches = ip == current_ip
            if "/" in ip:
                try:
                    current_matches = ipaddress.ip_address(current_ip) in ipaddress.ip_network(ip, strict=False)
                except ValueError:
                    current_matches = False
            if current_matches:
                return json_response(self, {"error": "Cannot block your current IP"}, HTTPStatus.BAD_REQUEST)
            if is_loopback_ip(ip):
                return json_response(self, {"error": "Cannot block localhost"}, HTTPStatus.BAD_REQUEST)
            with connect_db() as conn:
                conn.execute(
                    "INSERT OR REPLACE INTO ip_bans(ip, reason, banned_at, expires_at, banned_by) VALUES (?, ?, ?, NULL, ?)",
                    (ip, reason, now_iso(), admin["id"]),
                )
            return json_response(self, {"ok": True})
        if path.startswith("/api/security/unban-ip/"):
            if not self.require_user("admin"):
                return
            ip = urllib.parse.unquote(path.rsplit("/", 1)[1])
            with connect_db() as conn:
                conn.execute("DELETE FROM ip_bans WHERE ip = ?", (ip,))
            return json_response(self, {"ok": True})
        if path == "/api/import/run":
            user = self.require_user("write")
            if not user:
                return
            payload = read_json(self)
            title = (payload.get("title") or f"Auto import {now_iso()}").strip()
            vmware_host = (payload.get("vmwareHost") or "172.20.20.106").strip()
            veeam_host = (payload.get("veeamHost") or "172.20.20.22").strip()
            vmware_user = (payload.get("vmwareUser") or "").strip()
            vmware_password = payload.get("vmwarePassword") or ""
            veeam_user = (payload.get("veeamUser") or "").strip()
            veeam_password = payload.get("veeamPassword") or ""
            if not vmware_user or not vmware_password or not veeam_user or not veeam_password:
                return json_response(self, {"error": "VMware and Veeam credentials are required"}, HTTPStatus.BAD_REQUEST)
            try:
                write_import_log(user["id"], "running", f"Import started: VMware {vmware_host}, Veeam {veeam_host}")
                snapshot_id, summary = import_pair_from_servers(
                    title,
                    vmware_host,
                    vmware_user,
                    vmware_password,
                    veeam_host,
                    veeam_user,
                    veeam_password,
                    user["id"],
                )
            except subprocess.TimeoutExpired:
                write_import_log(user["id"], "error", "Import timed out")
                return json_response(self, {"error": "Import timed out"}, HTTPStatus.REQUEST_TIMEOUT)
            except Exception as exc:
                message = str(exc)
                write_import_log(user["id"], "error", "Import failed", message)
                return json_response(self, {"error": f"Import failed: {message}"}, HTTPStatus.BAD_REQUEST)
            write_import_log(user["id"], "success", f"Snapshot #{snapshot_id} created")
            return json_response(self, {"ok": True, "snapshotId": snapshot_id, "summary": summary}, HTTPStatus.CREATED)
        if path == "/api/export-filtered":
            if not self.require_user("read"):
                return
            payload = read_json(self)
            records = payload.get("records", [])
            snapshot = payload.get("snapshot", {})
            snapshot["title"] = f"{snapshot.get('title', 'Snapshot')} - filtered"
            snapshot["total_vms"] = len(records)
            snapshot["powered_on"] = sum(1 for item in records if item.get("powerState") == "PoweredOn")
            snapshot["powered_off"] = sum(1 for item in records if item.get("powerState") == "PoweredOff")
            snapshot["has_backup"] = sum(1 for item in records if item.get("hasVeeamBackup"))
            snapshot["has_replica"] = sum(1 for item in records if item.get("hasVeeamReplica"))
            snapshot["off_with_backup"] = sum(1 for item in records if item.get("powerState") == "PoweredOff" and item.get("hasVeeamBackup"))
            snapshot["off_with_replica"] = sum(1 for item in records if item.get("powerState") == "PoweredOff" and item.get("hasVeeamReplica"))
            output = BytesIO()
            write_report_xlsx(records, snapshot, output)
            body = output.getvalue()
            self.send_response(HTTPStatus.OK)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", 'attachment; filename="vsphere-veeam-filtered.xlsx"')
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        json_response(self, {"error": "Not found"}, HTTPStatus.NOT_FOUND)

    def handle_api_delete(self, path: str) -> None:
        if path.startswith("/api/users/"):
            user = self.require_user("admin")
            if not user:
                return
            try:
                user_id = int(path.rsplit("/", 1)[1])
            except ValueError:
                return json_response(self, {"error": "Invalid user id"}, HTTPStatus.BAD_REQUEST)
            if user_id == user["id"]:
                return json_response(self, {"error": "Cannot delete current user"}, HTTPStatus.BAD_REQUEST)
            with connect_db() as conn:
                if is_last_active_admin(conn, user_id):
                    return json_response(self, {"error": "Cannot delete the last active admin"}, HTTPStatus.BAD_REQUEST)
                cur = conn.execute("DELETE FROM users WHERE id = ?", (user_id,))
                if cur.rowcount == 0:
                    return json_response(self, {"error": "User not found"}, HTTPStatus.NOT_FOUND)
                conn.execute("DELETE FROM sessions WHERE user_id = ?", (user_id,))
            return json_response(self, {"ok": True})
        if path.startswith("/api/snapshots/"):
            if not self.require_user("write"):
                return
            snapshot_id = path.rsplit("/", 1)[1]
            if not snapshot_id.isdigit():
                return json_response(self, {"error": "Invalid snapshot id"}, HTTPStatus.BAD_REQUEST)
            target = safe_snapshot_dir(snapshot_id)
            if not target:
                return json_response(self, {"error": "Invalid snapshot path"}, HTTPStatus.BAD_REQUEST)
            with connect_db() as conn:
                cur = conn.execute("DELETE FROM snapshots WHERE id = ?", (snapshot_id,))
                if cur.rowcount == 0:
                    return json_response(self, {"error": "Snapshot not found"}, HTTPStatus.NOT_FOUND)
            if target.exists():
                shutil.rmtree(target)
            return json_response(self, {"ok": True})
        json_response(self, {"error": "Not found"}, HTTPStatus.NOT_FOUND)

    def handle_upload(self, user: sqlite3.Row) -> None:
        try:
            fields, files = parse_multipart_form(self)
        except ValueError as exc:
            return json_response(self, {"error": str(exc)}, HTTPStatus.BAD_REQUEST)

        title = (fields.get("title") or f"Snapshot {now_iso()}").strip()
        vsphere_file = files.get("vsphere")
        veeam_file = files.get("veeam")
        if not vsphere_file or not veeam_file:
            return json_response(self, {"error": "Both vSphere and Veeam CSV files are required"}, HTTPStatus.BAD_REQUEST)

        tmp_id = secrets.token_hex(8)
        tmp_dir = SNAPSHOT_DIR / f"tmp-{tmp_id}"
        tmp_dir.mkdir(parents=True, exist_ok=True)
        vsphere_path = tmp_dir / "vsphere.csv"
        veeam_path = tmp_dir / "veeam.csv"
        vsphere_name, vsphere_bytes = vsphere_file
        veeam_name, veeam_bytes = veeam_file
        vsphere_path.write_bytes(vsphere_bytes)
        veeam_path.write_bytes(veeam_bytes)

        try:
            records, summary = analyze_pair(vsphere_path, veeam_path)
        except Exception as exc:
            shutil.rmtree(tmp_dir, ignore_errors=True)
            return json_response(self, {"error": f"Failed to parse files: {exc}"}, HTTPStatus.BAD_REQUEST)

        with connect_db() as conn:
            cur = conn.execute(
                """
                INSERT INTO snapshots(
                    title, created_at, uploaded_by, vsphere_filename, veeam_filename,
                    total_vms, powered_on, powered_off, has_backup, has_replica, off_with_backup, off_with_replica
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    title,
                    now_iso(),
                    user["id"],
                    vsphere_name,
                    veeam_name,
                    summary["totalVms"],
                    summary["poweredOn"],
                    summary["poweredOff"],
                    summary["hasBackup"],
                    summary["hasReplica"],
                    summary["offWithBackup"],
                    summary["offWithReplica"],
                ),
            )
            snapshot_id = cur.lastrowid

        final_dir = SNAPSHOT_DIR / str(snapshot_id)
        tmp_dir.rename(final_dir)
        (final_dir / "analysis.json").write_text(json.dumps({"records": records, "summary": summary}, ensure_ascii=False, indent=2), encoding="utf-8")
        json_response(self, {"ok": True, "snapshotId": snapshot_id, "summary": summary}, HTTPStatus.CREATED)

    def load_snapshot(self, snapshot_id: str) -> tuple[dict, dict] | None:
        with connect_db() as conn:
            snapshot = conn.execute("SELECT * FROM snapshots WHERE id = ?", (snapshot_id,)).fetchone()
        if not snapshot:
            return None
        data_path = SNAPSHOT_DIR / str(snapshot_id) / "analysis.json"
        if not data_path.exists():
            return None
        return dict(snapshot), json.loads(data_path.read_text(encoding="utf-8"))

    def send_snapshot_data(self, snapshot_id: str) -> None:
        loaded = self.load_snapshot(snapshot_id)
        if not loaded:
            return json_response(self, {"error": "Snapshot not found"}, HTTPStatus.NOT_FOUND)
        snapshot, data = loaded
        json_response(self, {"snapshot": snapshot, **data})

    def send_export(self, snapshot_id: str) -> None:
        loaded = self.load_snapshot(snapshot_id)
        if not loaded:
            return json_response(self, {"error": "Snapshot not found"}, HTTPStatus.NOT_FOUND)
        snapshot, data = loaded
        output = BytesIO()
        write_report_xlsx(data["records"], snapshot, output)
        body = output.getvalue()
        filename = f"vsphere-veeam-{snapshot_id}.xlsx"
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)


def main() -> None:
    init_db()
    host = os.environ.get("VM_PORTAL_HOST", "127.0.0.1")
    port = int(os.environ.get("VM_PORTAL_PORT", "8088"))
    server = ThreadingHTTPServer((host, port), PortalHandler)
    print(f"VM portal running at http://{host}:{port}")
    print("Default admin user: codex")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass


if __name__ == "__main__":
    main()

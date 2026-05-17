from __future__ import annotations

import csv
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


VSPHERE_PATH = Path(r"C:\Users\vaisman\Desktop\vsphere-vm-power-state.csv")
VEEAM_PATH = Path(r"C:\Users\vaisman\Desktop\veeam-backups-and-replicas.csv")
OUTPUT_PATH = Path(r"C:\Users\vaisman\Desktop\vsphere-veeam-full-report.xlsx")


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [{k: (v or "").strip() for k, v in row.items()} for row in csv.DictReader(handle)]


def norm(value: str) -> str:
    return value.strip().casefold()


def parse_date(value: str) -> datetime | None:
    try:
        return datetime.strptime(value, "%d/%m/%Y %H:%M:%S")
    except ValueError:
        return None


vsphere_rows = read_csv(VSPHERE_PATH)
veeam_rows = read_csv(VEEAM_PATH)

veeam_by_vm: dict[str, dict[str, dict[str, object]]] = defaultdict(
    lambda: {
        "Backup": {
            "jobs": set(),
            "count": 0,
            "latest": None,
            "latest_text": "",
        },
        "Replica": {
            "jobs": set(),
            "count": 0,
            "latest": None,
            "latest_text": "",
        },
    }
)

for row in veeam_rows:
    vm_name = row.get("VMName", "")
    protection = row.get("Protection", "")
    if not vm_name or protection not in {"Backup", "Replica"}:
        continue
    item = veeam_by_vm[norm(vm_name)][protection]
    item["jobs"].add(row.get("JobName", ""))
    item["count"] += 1
    created = parse_date(row.get("CreationTime", ""))
    if created and (item["latest"] is None or created > item["latest"]):
        item["latest"] = created
        item["latest_text"] = row.get("CreationTime", "")

report_rows = []
for row in vsphere_rows:
    vm_name = row.get("Name", "")
    vm_info = veeam_by_vm.get(norm(vm_name))
    backup = vm_info["Backup"] if vm_info else None
    replica = vm_info["Replica"] if vm_info else None
    backup_count = int(backup["count"]) if backup else 0
    replica_count = int(replica["count"]) if replica else 0
    report_rows.append(
        {
            "VMName": vm_name,
            "PowerState": row.get("PowerState", ""),
            "VMHost": row.get("VMHost", ""),
            "IsReplicaVM": "Yes" if "replica" in norm(vm_name) else "No",
            "HasVeeamBackup": "Yes" if backup_count else "No",
            "HasVeeamReplica": "Yes" if replica_count else "No",
            "HasAnyVeeamProtection": "Yes" if backup_count or replica_count else "No",
            "BackupJobs": "; ".join(sorted(job for job in (backup or {}).get("jobs", set()) if job)),
            "ReplicaJobs": "; ".join(sorted(job for job in (replica or {}).get("jobs", set()) if job)),
            "LastBackupRestorePoint": (backup or {}).get("latest_text", ""),
            "LastReplicaRestorePoint": (replica or {}).get("latest_text", ""),
            "BackupRestorePoints": backup_count or "",
            "ReplicaRestorePoints": replica_count or "",
        }
    )

report_rows.sort(key=lambda item: (item["PowerState"], item["IsReplicaVM"], item["VMName"].casefold()))

wb = Workbook()
ws = wb.active
ws.title = "All VMs"

headers = [
    "VMName",
    "PowerState",
    "VMHost",
    "IsReplicaVM",
    "HasVeeamBackup",
    "HasVeeamReplica",
    "HasAnyVeeamProtection",
    "BackupJobs",
    "ReplicaJobs",
    "LastBackupRestorePoint",
    "LastReplicaRestorePoint",
    "BackupRestorePoints",
    "ReplicaRestorePoints",
]
ws.append(headers)
for item in report_rows:
    ws.append([item[header] for header in headers])

header_fill = PatternFill("solid", fgColor="1F2937")
backup_fill = PatternFill("solid", fgColor="F8BBD0")
replica_fill = PatternFill("solid", fgColor="D9EAD3")
both_fill = PatternFill("solid", fgColor="D9B8FF")
off_fill = PatternFill("solid", fgColor="F3F4F6")
white_font = Font(color="FFFFFF", bold=True)
thin_gray = Side(style="thin", color="D1D5DB")
border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = white_font
    cell.border = border
    cell.alignment = Alignment(horizontal="center")

for row_idx in range(2, ws.max_row + 1):
    power_state = ws.cell(row_idx, 2).value
    has_backup = ws.cell(row_idx, 5).value == "Yes"
    has_replica = ws.cell(row_idx, 6).value == "Yes"
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row_idx, col_idx)
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=col_idx in {8, 9})
    if power_state == "PoweredOff":
        ws.cell(row_idx, 2).fill = off_fill
    if has_backup and has_replica:
        ws.cell(row_idx, 1).fill = both_fill
        ws.cell(row_idx, 5).fill = both_fill
        ws.cell(row_idx, 6).fill = both_fill
    elif has_backup:
        ws.cell(row_idx, 1).fill = backup_fill
        ws.cell(row_idx, 5).fill = backup_fill
    elif has_replica:
        ws.cell(row_idx, 1).fill = replica_fill
        ws.cell(row_idx, 6).fill = replica_fill

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
widths = {
    "A": 42,
    "B": 14,
    "C": 18,
    "D": 12,
    "E": 16,
    "F": 16,
    "G": 20,
    "H": 32,
    "I": 32,
    "J": 22,
    "K": 22,
    "L": 18,
    "M": 18,
}
for col, width in widths.items():
    ws.column_dimensions[col].width = width

summary = wb.create_sheet("Summary")


def count_where(**criteria: str) -> int:
    total = 0
    for item in report_rows:
        if all(item[key] == value for key, value in criteria.items()):
            total += 1
    return total


summary_rows = [
    ["Metric", "Count"],
    ["All VMs", len(report_rows)],
    ["PoweredOn", count_where(PowerState="PoweredOn")],
    ["PoweredOff", count_where(PowerState="PoweredOff")],
    ["Replica VMs in vSphere name", count_where(IsReplicaVM="Yes")],
    ["Has Veeam Backup", count_where(HasVeeamBackup="Yes")],
    ["Has Veeam Replica", count_where(HasVeeamReplica="Yes")],
    ["Has Backup and Replica", sum(1 for item in report_rows if item["HasVeeamBackup"] == "Yes" and item["HasVeeamReplica"] == "Yes")],
    ["No Veeam protection found", count_where(HasAnyVeeamProtection="No")],
    ["PoweredOff with Veeam Backup", count_where(PowerState="PoweredOff", HasVeeamBackup="Yes")],
    ["PoweredOff with Veeam Replica", count_where(PowerState="PoweredOff", HasVeeamReplica="Yes")],
    ["PoweredOff without Veeam protection", count_where(PowerState="PoweredOff", HasAnyVeeamProtection="No")],
]
for row in summary_rows:
    summary.append(row)
for cell in summary[1]:
    cell.fill = header_fill
    cell.font = white_font
    cell.border = border
for row in summary.iter_rows(min_row=2, max_row=summary.max_row, max_col=2):
    for cell in row:
        cell.border = border
summary.column_dimensions["A"].width = 42
summary.column_dimensions["B"].width = 14

legend = wb.create_sheet("Legend")
legend_rows = [
    ["Color", "Meaning"],
    ["Pink", "VM has Veeam Backup restore points"],
    ["Green", "VM has Veeam Replica restore points"],
    ["Purple", "VM has both Backup and Replica restore points"],
    ["Gray in PowerState", "VM is PoweredOff"],
]
for row in legend_rows:
    legend.append(row)
for cell in legend[1]:
    cell.fill = header_fill
    cell.font = white_font
    cell.border = border
legend["A2"].fill = backup_fill
legend["A3"].fill = replica_fill
legend["A4"].fill = both_fill
legend["A5"].fill = off_fill
for row in legend.iter_rows(min_row=2, max_row=legend.max_row, max_col=2):
    for cell in row:
        cell.border = border
legend.column_dimensions["A"].width = 16
legend.column_dimensions["B"].width = 46

wb.save(OUTPUT_PATH)

check_wb = load_workbook(OUTPUT_PATH, read_only=False)
check_ws = check_wb["All VMs"]
print(f"OUTPUT={OUTPUT_PATH}")
print(f"ROWS={check_ws.max_row - 1}")
print(f"COLS={check_ws.max_column}")
print(f"POWERED_ON={count_where(PowerState='PoweredOn')}")
print(f"POWERED_OFF={count_where(PowerState='PoweredOff')}")
print(f"HAS_BACKUP={count_where(HasVeeamBackup='Yes')}")
print(f"HAS_REPLICA={count_where(HasVeeamReplica='Yes')}")
print(f"OFF_WITH_BACKUP={count_where(PowerState='PoweredOff', HasVeeamBackup='Yes')}")
print(f"OFF_WITH_REPLICA={count_where(PowerState='PoweredOff', HasVeeamReplica='Yes')}")

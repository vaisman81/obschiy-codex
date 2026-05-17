from __future__ import annotations

import csv
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


VSPHERE_PATH = Path(r"C:\Users\vaisman\Desktop\vsphere-vm-power-state.csv")
VEEAM_PATH = Path(r"C:\Users\vaisman\Desktop\veeam-restore-points.csv")
OUTPUT_PATH = Path(r"C:\Users\vaisman\Desktop\poweredoff-servers-veeam-marked-fixed.xlsx")


def clean_line(line: str) -> str:
    line = line.rstrip("\r\n").lstrip("\ufeff")
    if line.startswith('"'):
        line = line[1:]
    if line.endswith('"'):
        line = line[:-1]
    return line.replace('""', '"')


def read_broken_csv(path: Path) -> list[dict[str, str]]:
    lines = [clean_line(line) for line in path.read_text(encoding="utf-8-sig").splitlines() if line.strip()]
    reader = csv.DictReader(lines)
    return [{k: (v or "").strip() for k, v in row.items()} for row in reader]


def norm(value: str) -> str:
    return value.strip().casefold()


def parse_veeam_date(value: str) -> datetime | None:
    try:
        return datetime.strptime(value, "%d/%m/%Y %H:%M:%S")
    except ValueError:
        return None


vsphere_rows = read_broken_csv(VSPHERE_PATH)
veeam_rows = read_broken_csv(VEEAM_PATH)

veeam_by_vm: dict[str, dict[str, object]] = {}
for row in veeam_rows:
    vm_name = row.get("VMName", "").strip()
    if not vm_name:
        continue
    key = norm(vm_name)
    item = veeam_by_vm.setdefault(
        key,
        {
            "vm_name": vm_name,
            "backup_names": set(),
            "restore_point_count": 0,
            "last_restore_point": None,
            "last_restore_point_text": "",
        },
    )
    item["backup_names"].add(row.get("BackupName", "").strip())
    item["restore_point_count"] += 1
    parsed_date = parse_veeam_date(row.get("CreationTime", ""))
    if parsed_date and (item["last_restore_point"] is None or parsed_date > item["last_restore_point"]):
        item["last_restore_point"] = parsed_date
        item["last_restore_point_text"] = row.get("CreationTime", "")

report_rows = []
for row in vsphere_rows:
    vm_name = row.get("Name", "").strip()
    if row.get("PowerState") != "PoweredOff":
        continue
    if "replica" in norm(vm_name):
        continue

    veeam = veeam_by_vm.get(norm(vm_name))
    report_rows.append(
        {
            "server": vm_name,
            "power_state": row.get("PowerState", ""),
            "veeam_backup": "Yes" if veeam else "",
            "backup_jobs": "; ".join(sorted(name for name in (veeam or {}).get("backup_names", set()) if name)),
            "last_restore_point": (veeam or {}).get("last_restore_point_text", ""),
            "restore_points": (veeam or {}).get("restore_point_count", ""),
        }
    )

report_rows.sort(key=lambda item: item["server"].casefold())

wb = Workbook()
ws = wb.active
ws.title = "PoweredOff Servers"

headers = ["Server", "Power State", "Veeam Backup", "Backup Job(s)", "Last Restore Point", "Restore Points"]
ws.append(headers)
for item in report_rows:
    ws.append(
        [
            item["server"],
            item["power_state"],
            item["veeam_backup"],
            item["backup_jobs"],
            item["last_restore_point"],
            item["restore_points"],
        ]
    )

header_fill = PatternFill("solid", fgColor="1F2937")
pink_fill = PatternFill("solid", fgColor="F8BBD0")
white_font = Font(color="FFFFFF", bold=True)
thin_gray = Side(style="thin", color="D1D5DB")
border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = white_font
    cell.border = border
    cell.alignment = Alignment(horizontal="center")

for row_idx in range(2, ws.max_row + 1):
    marked = ws.cell(row_idx, 3).value == "Yes"
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row_idx, col_idx)
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=(col_idx == 4))
    if marked:
        ws.cell(row_idx, 1).fill = pink_fill
        ws.cell(row_idx, 3).fill = pink_fill

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
ws.column_dimensions["A"].width = 42
ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 14
ws.column_dimensions["D"].width = 36
ws.column_dimensions["E"].width = 20
ws.column_dimensions["F"].width = 14

summary = wb.create_sheet("Summary")
summary_rows = [
    ["Metric", "Count"],
    ["Powered off servers excluding replicas", len(report_rows)],
    ["Pink marked - found in Veeam", sum(1 for row in report_rows if row["veeam_backup"] == "Yes")],
    ["Not marked - not found in Veeam", sum(1 for row in report_rows if row["veeam_backup"] != "Yes")],
]
for row in summary_rows:
    summary.append(row)
for cell in summary[1]:
    cell.fill = header_fill
    cell.font = white_font
    cell.border = border
for row in summary.iter_rows(min_row=2, max_row=summary.max_row, max_col=2):
    for cell in row:
        cell.border = border
summary.column_dimensions["A"].width = 42
summary.column_dimensions["B"].width = 12

wb.save(OUTPUT_PATH)

check_wb = load_workbook(OUTPUT_PATH)
check_ws = check_wb["PoweredOff Servers"]
pink_count = sum(1 for row in range(2, check_ws.max_row + 1) if check_ws.cell(row, 1).fill.fgColor.rgb in {"00F8BBD0", "F8BBD0"})
print(f"OUTPUT={OUTPUT_PATH}")
print(f"ROWS={check_ws.max_row - 1}")
print(f"PINK_MARKED={pink_count}")
print(f"REPLICAS_IN_REPORT={sum(1 for row in range(2, check_ws.max_row + 1) if 'replica' in str(check_ws.cell(row, 1).value).casefold())}")
